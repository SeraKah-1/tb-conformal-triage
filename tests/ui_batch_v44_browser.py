"""Batch controls, real folder input, XLSX round-trip, and reset isolation. No model inference."""
import base64,json,os,socket,subprocess,tempfile,time,urllib.request,zipfile
from pathlib import Path
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
checks=[]
def ck(name,ok):
    assert ok,name
    checks.append(name)
with socket.socket() as sock:sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
server=subprocess.Popen(['node','tb_pwa_offline_workstation/server.js'],cwd=ROOT,env={**os.environ,'PORT':str(port)},stdout=subprocess.DEVNULL)
png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jUVkAAAAASUVORK5CYII=')
try:
    for _ in range(100):
        try:urllib.request.urlopen(f'http://127.0.0.1:{port}',timeout=.2).close();break
        except OSError:time.sleep(.05)
    with tempfile.TemporaryDirectory() as tmp,sync_playwright() as pw:
        folder=Path(tmp)/'images';(folder/'nested').mkdir(parents=True)
        for i in range(1000):(folder/('nested' if i%2 else '')/f'image-{i}.png').write_bytes(png)
        (folder/'notes.txt').write_text('Not an image')
        browser=pw.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox']);context=browser.new_context(service_workers='block',accept_downloads=True);page=context.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
        page.add_init_script("localStorage.setItem('tb_triage_disclaimer_accepted','true')");page.goto(f'http://127.0.0.1:{port}');page.wait_for_function('window.caseWorkspace && !caseWorkspace.state.loadingArchives');page.locator('[data-case-view=queue]').click()
        ck('Back is visible in Batch',page.locator('#case-workspace').get_by_role('button',name='Kembali',exact=True).is_visible())
        ck('only one Process action before selection',page.locator('#case-process').is_visible() and not page.locator('#case-process-selected').is_visible())
        page.locator('#case-folder').set_input_files(str(folder));page.wait_for_function('caseWorkspace.state.records.length===1000')
        ck('folder imports 1000 images in one selection',page.evaluate('caseWorkspace.state.records.length===1000'))
        ck('nested paths retained',page.evaluate('caseWorkspace.state.records.some(r=>r.sourcePath.includes("nested/"))'))
        ck('unsupported file reported', 'notes.txt' in page.locator('#case-workspace .case-message').inner_text())
        page.screenshot(path='/workspace/audit-tb/deployment-v4.4/batch-local.png',full_page=True)
        ck('1000 inputs bounded to 50 rows',page.locator('.case-row').count()==50)
        page.locator('.case-more > summary').click();page.locator('#case-select-all').click();ck('select all spans every page',page.evaluate('caseWorkspace.state.selectedIds.length===1000'));ck('one Process action after selection',page.locator('#case-process-selected').is_visible() and not page.locator('#case-process').is_visible())
        page.evaluate(r'''() => {const r=caseWorkspace.state.records[0];r.filename='=SUM(1,2).png';r.notes='Baris satu, koma\nBaris dua: α & <teks>';r.state='complete';r.result={filename:r.filename,triage_action:'AUTO_FLAG_TB_URGENT',probability_tb:.9,probability_normal:.1,conformal_details:{conformal_set:['TB']},iqa_report:{quality_passed:true},ood_report:{available:true,is_ood:false,mahalanobis_distance:1,threshold:3},hirescam_heatmap_base64:'data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jUVkAAAAASUVORK5CYII=',allow_autonomous_release:false};caseWorkspace.setLanguage('id');}''')
        with page.expect_download() as down:page.locator('#case-export').click()
        ck('Excel filename',down.value.suggested_filename.endswith('.xlsx'));book=load_workbook(open(down.value.path(),'rb'));sheet=book.active
        ck('Excel has 1000 rows and consistent columns',sheet.max_row==1001 and sheet.max_column==16)
        values=list(sheet.values);idx=next(i for i,r in enumerate(values) if r[1]=='=SUM(1,2).png');row=values[idx]
        ck('formula-like filename stored as text',sheet.cell(idx+1,2).data_type=='s')
        ck('Unicode multiline notes intact',row[13]=='Baris satu, koma\nBaris dua: α & <teks>')
        ck('score formatted as percent',row[8]==.9 and sheet.cell(idx+1,9).number_format=='0.00%');ck('freeze header and filter',sheet.freeze_panes=='A2' and sheet.auto_filter.ref=='A1:P1001')
        ck('zip CRC valid',zipfile.ZipFile(down.value.path()).testzip() is None)
        page.locator('[data-heatmap-id]').click();ck('Heatmap opens corresponding result',page.evaluate('caseWorkspace.state.activeId===caseWorkspace.state.records[0].id') and page.locator('#results-display').is_visible());ck('heatmap uses stored image',page.locator('#img-heat').get_attribute('src').startswith('data:image/png'));ck('IQA pass is green',page.locator('[data-mechanism=IQA]').get_attribute('data-tone')=='success')
        page.locator('#case-detail-tools').get_by_role('button',name='Kembali',exact=True).click();ck('Back returns to list and retains selection',page.evaluate('caseWorkspace.state.activeId===null&&caseWorkspace.state.selectedIds.length===1000'))
        page.evaluate('caseWorkspace.open(caseWorkspace.state.records[0].id)');page.evaluate('caseWorkspace.save()');page.wait_for_function('caseWorkspace.state.records[0].saved&&!caseWorkspace.state.records[0].saving')
        page.evaluate('caseWorkspace.state.records[0].result.hirescam_heatmap_base64=null;caseWorkspace.open(caseWorkspace.state.records[0].id)');ck('missing heatmap is explicit and controls disabled',page.locator('#case-heatmap-notice').is_visible() and page.locator('#btn-mode-heat').is_disabled() and page.locator('#btn-mode-orig').evaluate('(el)=>el.classList.contains("active")'))
        page.once('dialog',lambda d:d.dismiss());page.evaluate('caseWorkspace.resetSession()');ck('cancel reset preserves session',page.evaluate('caseWorkspace.state.records.length===1000'))
        page.once('dialog',lambda d:d.accept());page.evaluate('caseWorkspace.resetSession()');ck('reset clears session and selection',page.evaluate('caseWorkspace.state.records.length===0&&caseWorkspace.state.selectedIds.length===0&&currentResultData===null'))
        page.locator('[data-case-view=saved]').click();page.wait_for_function('caseWorkspace.state.records.length===1&&!caseWorkspace.state.loadingArchives');ck('reset preserves saved case on disk',page.evaluate('caseWorkspace.state.records[0].saved'))
        page.locator('[data-case-view=queue]').click()
        await_drop='''async () => {const file=new File(['png'],'dropped.png',{type:'image/png'});let calls=0;const leaf={isFile:true,name:file.name,file:cb=>cb(file)},dir={isDirectory:true,name:'dropped',createReader:()=>({readEntries:cb=>cb(calls++===0?[leaf]:[])})};await document.getElementById('case-workspace').ondrop({preventDefault(){},dataTransfer:{items:[{webkitGetAsEntry:()=>dir}],files:[]}});}'''
        page.evaluate(await_drop);ck('drop reads directory entries to exhaustion',page.evaluate('caseWorkspace.state.records.some(r=>r.sourcePath==="dropped/dropped.png")'))
        before=page.evaluate('caseWorkspace.state.records.length');page.evaluate('''async () => {let calls=0;const dir={isDirectory:true,name:'bad',createReader:()=>({readEntries:(ok,bad)=>{if(calls++===0)ok([{isFile:true,file:cb=>cb(new File(['x'],'partial.png'))}]);else bad(new Error('denied'));}})};await document.getElementById('case-workspace').ondrop({preventDefault(){},dataTransfer:{items:[{webkitGetAsEntry:()=>dir}],files:[]}});}''');ck('folder read failure adds no partial files',page.evaluate('caseWorkspace.state.records.length')==before)
        page.evaluate('''() => {processBatchItem=()=>new Promise(resolve=>window.finishJob=()=>resolve({triage_action:'ERROR_RUNTIME'}));window.job=caseWorkspace.process();}''');page.evaluate('caseWorkspace.resetSession()');ck('reset blocked during active processing',page.evaluate('caseWorkspace.state.running&&caseWorkspace.state.records.length===2'));page.wait_for_function('typeof finishJob === "function"');page.evaluate('finishJob()');page.wait_for_function('!caseWorkspace.state.running')
        page.evaluate("caseWorkspace.state.records[0].notes='x'.repeat(32768)");page.evaluate('caseWorkspace.export()');ck('Excel limit reported without silent truncation','CSV' in page.locator('#case-workspace .case-message').inner_text())
        ck('no JS errors',not errors);context.close();browser.close()
    print(json.dumps({'passed':len(checks),'checks':checks},indent=2))
finally:server.terminate();server.wait()
